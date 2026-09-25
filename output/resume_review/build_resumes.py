from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent

def setup(cn=False):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin = s.bottom_margin = Inches(.48)
    s.left_margin = s.right_margin = Inches(.58)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = d.styles[name]
        st.font.name = 'Calibri'
        st.font.color.rgb = RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    n = d.styles['Normal']
    n.font.size = Pt(10 if cn else 10.5)
    n.paragraph_format.space_after = Pt(2.5)
    n.paragraph_format.line_spacing = 1.04
    n.paragraph_format.widow_control = True
    h = d.styles['Heading 1']
    h.font.size = Pt(11.5)
    h.font.bold = True
    h.paragraph_format.space_before = Pt(7)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    t = d.styles['Title']
    t.font.size = Pt(19)
    t.paragraph_format.space_after = Pt(2)
    d.core_properties.author = 'Chihan Gao'
    d.core_properties.title = 'Chihan Gao Quantitative Research Internship Resume'
    p = d.add_paragraph('Chihan (Grant) Gao', 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph('grantgao2002@outlook.com  |  +1 425-503-5652  |  linkedin.com/in/chihangao')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.font.size=Pt(9)
    p=d.add_paragraph('求职方向  量化研究与资金管理实习' if cn else 'Quantitative Research and Portfolio Management Internship')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    return d

def head(d, text):
    d.add_paragraph(text, 'Heading 1')

def line(d, title, date='', sub=False):
    p=d.add_paragraph()
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(1.5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.11), WD_TAB_ALIGNMENT.RIGHT)
    r=p.add_run(title); r.bold=not sub
    if date: p.add_run('\t'+date)
    return p

def bullet(d, text):
    p=d.add_paragraph()
    p.paragraph_format.left_indent=Inches(.12)
    p.paragraph_format.first_line_indent=Inches(-.12)
    p.add_run('•  '+text)

def skill(d, label, text):
    p=d.add_paragraph()
    p.add_run(label).bold=True
    p.add_run(text)

def finish(d, path, cn=False):
    for st in d.styles:
        for node in list(st.element.iter(qn('w:pBdr'))):
            node.getparent().remove(node)
        if st.type == 1:
            pr=st.element.get_or_add_pPr()
            snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pr.append(snap)
    for p in d.paragraphs:
        pr=p._p.get_or_add_pPr()
        for node in list(pr.findall(qn('w:pBdr'))): pr.remove(node)
        snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pr.append(snap)
        if cn and p.style.name not in ['Title', 'Heading 1']:
            p.paragraph_format.line_spacing=Pt(14)
    for node in list(d.settings.element.findall(qn('w:updateStyles'))):
        d.settings.element.remove(node)
    d.save(path)

en=setup()
head(en,'Education')
line(en,'WorldQuant University','Apr 2026 – Mar 2028 (expected)')
line(en,'MSc in Financial Engineering | Coursework: Financial Data; Financial Markets',sub=True)
line(en,'University of Washington, Seattle','Sep 2019 – Dec 2022')
line(en,'BS in Real Estate | GPA: 3.63/4.0 | Dean’s List: 7 times',sub=True)
en.add_paragraph('Relevant coursework: Financial Modeling; Data Modeling; Project Management')

head(en,'Technical Skills')
skill(en,'Programming and data: ','Python, pandas, NumPy, SciPy, Jupyter Notebook, Excel; data cleaning and EDA')
skill(en,'Quantitative research: ','Factor testing, feature selection, MLP classification, time-ordered validation, portfolio backtesting, drawdown and transaction-cost analysis')
skill(en,'Financial analysis: ','Cash-flow analysis, DCF, IRR, NOI forecasting, valuation and scenario analysis')

head(en,'Quantitative Projects')
line(en,'A Share Multi Factor Research and Portfolio Backtesting','Python | 2026')
bullet(en,'Built a modular research workflow covering market-data preparation, factor calculation, single-factor testing, composite signals, portfolio construction and performance reporting.')
bullet(en,'Implemented price, volume, volatility and liquidity factors with reusable registration and caching; evaluated factor IC/IR and correlations to support factor selection and weighting.')
bullet(en,'Implemented rolling validation with 24-month training and 6-month test windows; modeled monthly rebalancing and turnover costs, and examined portfolio returns and drawdowns across test periods.')

line(en,'Market Direction Classification and Alternative Data','Python | 2026')
bullet(en,'Replicated an ECH market-direction study using 2,768 daily observations and 20 price/volume features; compared full and reduced feature sets using Pearson selection and a scaled MLP.')
bullet(en,'Re-evaluated the model using next-day targets and time-ordered validation, obtaining approximately 52–53% accuracy versus approximately 80% under the original evaluation setup; examined information timing and validation assumptions.')
bullet(en,'Processed 100 Google News metadata records from 45 sources, checking timestamps, duplicates and source concentration to assess alternative-data quality.')

line(en,'SPY and IWM Portfolio Rebalancing','Python | 2026')
bullet(en,'Built a backtest of a 50/50 SPY–IWM portfolio using approximately 10 years of daily data, tracking shares, asset values, cash balances and portfolio-weight drift.')
bullet(en,'Implemented buy/sell and post-trade reconciliation logic; analyzed cumulative performance, drawdowns and rebalancing requirements.')

head(en,'Professional Experience')
line(en,'Vibrant Cities, LLC | Seattle','Jan 2023 – Present')
line(en,'Financial / Asset Manager','Nov 2024 – Present',sub=True)
en.add_paragraph('Promoted from Assistant Asset Manager and Development Analyst')
bullet(en,'Analyze financial statements, cash flows, rent rolls and budget variances across a 350+ unit real-asset portfolio to identify revenue, expense and NOI trends.')
bullet(en,'Perform DCF valuation and scenario analysis incorporating NOI growth, debt assumptions, cap rates and exit values to support investment and operating decisions.')
bullet(en,'Evaluate redevelopment opportunities in ROIC’s REIT portfolio using public and geospatial data; support asset dispositions and due diligence with financial packages and property-level analysis.')
line(en,'Binarix | Shanghai','Jun 2021 – Sep 2021')
line(en,'Data Analyst Summer Intern',sub=True)
bullet(en,'Built recurring analytics for 10–15 products, tracking exposure, click-through, conversion and sales; presented findings to support advertising and customer-targeting decisions.')
finish(en, OUT/'Chihan_Gao_Quant_Resume_EN.docx')

cn=setup(True)
head(cn,'教育背景')
line(cn,'WorldQuant University','2026.04 – 2028.03（预计）')
line(cn,'金融工程硕士在读 | 相关课程：金融数据、金融市场',sub=True)
line(cn,'华盛顿大学西雅图分校 University of Washington','2019.09 – 2022.12')
line(cn,'房地产专业理学学士 | GPA：3.63/4.0 | 7 次入选 Dean’s List',sub=True)
cn.add_paragraph('相关课程：财务建模、数据建模、项目管理')

head(cn,'专业技能')
skill(cn,'编程与数据：','Python、pandas、NumPy、SciPy、Jupyter Notebook、Excel；数据清洗与探索性分析')
skill(cn,'量化研究：','因子检验、特征筛选、MLP 分类、时间顺序验证、组合回测、回撤与交易成本分析')
skill(cn,'财务分析：','现金流分析、DCF 估值、IRR 测算、净营业收入预测、估值与情景分析')

head(cn,'量化研究项目')
line(cn,'A 股多因子研究与组合回测','Python | 2026')
bullet(cn,'搭建模块化研究流程，覆盖行情数据准备、因子计算、单因子检验、信号合成、组合构建及绩效报告。')
bullet(cn,'实现量价、波动率与流动性等因子，设置可复用的因子注册与缓存机制；分析因子 IC、IR 及相关性，为因子筛选和权重配置提供依据。')
bullet(cn,'实现 24 个月训练、6 个月测试的滚动验证；在回测中设置月度调仓与换手成本，分析不同测试阶段的组合收益及回撤表现。')

line(cn,'市场涨跌分类与另类数据处理','Python | 2026')
bullet(cn,'复现 ECH 市场方向分类研究，使用 2,768 条日度观测构建 20 个量价特征，结合 Pearson 相关性筛选与标准化 MLP，对比完整及精简特征集。')
bullet(cn,'重新采用次日预测目标与时间顺序验证，准确率约为 52–53%，低于原评估设置下约 80% 的结果；分析信息可用时点与验证假设对结果的影响。')
bullet(cn,'处理来自 45 个来源的 100 条 Google News 元数据，检查时间戳、重复记录及来源集中度，评估另类数据质量。')

line(cn,'SPY 与 IWM 投资组合再平衡','Python | 2026')
bullet(cn,'基于约 10 年日度行情，搭建目标权重为 50/50 的 SPY–IWM 组合回测，跟踪持仓数量、资产市值、现金余额及权重偏移。')
bullet(cn,'实现买卖交易与交易后账目核对逻辑，分析组合累计表现、回撤及再平衡需求。')

head(cn,'工作与实习经历')
line(cn,'Vibrant Cities, LLC | 西雅图','2023.01 – 至今')
line(cn,'财务及资产经理','2024.11 – 至今',sub=True)
cn.add_paragraph('由助理资产经理及开发分析师晋升')
bullet(cn,'分析覆盖 350 余套房源的资产组合，结合财务报表、现金流、租赁台账及预算差异，识别收入、费用与净营业收入的变化。')
bullet(cn,'开展 DCF 估值与情景分析，结合净营业收入增长、债务假设、资本化率及退出价值，为投资与运营决策提供支持。')
bullet(cn,'利用公开资料及地理空间数据研究 ROIC 房地产信托资产组合的再开发机会；整理财务资料与资产经营分析，支持资产出售及尽职调查。')
line(cn,'Binarix | 上海','2021.06 – 2021.09')
line(cn,'数据分析暑期实习生',sub=True)
bullet(cn,'搭建覆盖约 10–15 款产品的定期分析流程，跟踪曝光、点击、转化及销售指标；向市场与管理团队汇报结果，支持广告投放及客群定位调整。')
finish(cn, OUT/'Chihan_Gao_Quant_Resume_CN.docx', True)
print('Created EN and CN resumes')
